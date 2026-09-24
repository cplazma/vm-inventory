import requests
import urllib3
import re
import json
import os
from datetime import datetime
import pandas as pd
from openpyxl.styles import Border, Side, Alignment, PatternFill

# Suppress SSL warnings for self-signed certificates
urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

def load_config(filepath="servers.json"):
    try:
        with open(filepath, 'r') as file:
            return json.load(file)
    except FileNotFoundError:
        print(f"Error: Configuration file '{filepath}' not found.")
        exit(1)
    except json.JSONDecodeError:
        print(f"Error: '{filepath}' is not valid JSON.")
        exit(1)

def load_vlan_zones(filepath="vlan-zone.txt"):
    """Load the VLAN to Zone mappings into a dictionary."""
    zones = {}
    try:
        with open(filepath, 'r') as file:
            lines = file.readlines()
            for line in lines[1:]: # Skip the header row
                parts = line.strip().split()
                if len(parts) >= 3:
                    node = parts[0].strip()
                    vlan = parts[1].strip().lower()
                    zone = parts[2].strip()
                    zones[(node, vlan)] = zone
    except FileNotFoundError:
        print(f"Warning: '{filepath}' not found. Zone mapping will be skipped.")
    return zones

def determine_zone(node, vlan_str, zone_map):
    """Map the extracted VLAN string to a Zone based on the lookup table."""
    if vlan_str == "None":
        return zone_map.get((node, "none"), zone_map.get(("ALL", "none"), "Unknown"))
    
    # Extract just the numbers from the network string (e.g., "net0(103)" -> "103")
    extracted_vlans = re.findall(r'\((\d+)\)', vlan_str)
    if not extracted_vlans:
        return "Unknown"
    
    mapped_zones = []
    for v_str in extracted_vlans:
        z = zone_map.get((node, v_str), zone_map.get(("ALL", v_str), "Unknown"))
        if z not in mapped_zones:
            mapped_zones.append(z)
            
    return ", ".join(mapped_zones)

def make_request(method, url, headers, silent_agent=False):
    try:
        response = requests.request(method, url, headers=headers, verify=False, timeout=10)
        response.raise_for_status()
        return response.json().get('data', None)
    except requests.exceptions.HTTPError as e:
        if silent_agent and e.response.status_code in [500, 501]: 
            return None
        print(f"  -> API Error: {e.response.status_code} - {e.response.text}")
        return None
    except requests.exceptions.RequestException as e:
        print(f"  -> Connection Error: {e}")
        return None

def extract_vlan(config):
    vlans = []
    for key, val in config.items():
        if key.startswith('net'):
            match = re.search(r'tag=(\d+)', str(val))
            if match: vlans.append(f"{key}({match.group(1)})")
    return ", ".join(vlans) if vlans else "None"

def extract_disks(config):
    disks = []
    total_gb = 0.0
    bus_prefixes = ('scsi', 'sata', 'ide', 'virtio', 'rootfs', 'mp')
    for key, val in config.items():
        if key.startswith(bus_prefixes) and 'media=cdrom' not in str(val):
            match = re.search(r'size=([0-9\.]+)([a-zA-Z]+)', str(val))
            if match:
                size_val = float(match.group(1))
                unit = match.group(2).upper()
                if unit == 'T': total_gb += size_val * 1024
                elif unit == 'G': total_gb += size_val
                elif unit == 'M': total_gb += size_val / 1024
                elif unit == 'K': total_gb += size_val / (1024 * 1024)
                disks.append(f"{key}:{match.group(1)}{match.group(2)}")
    return ", ".join(disks) if disks else "None", round(total_gb, 2)

def get_guest_os(base_url, headers, node, vmid):
    url = f"{base_url}/nodes/{node}/qemu/{vmid}/agent/get-os-info"
    agent_data = make_request("GET", url, headers, silent_agent=True)
    if agent_data and 'result' in agent_data: return agent_data['result'].get('pretty-name')
    return None

def get_guest_ips(base_url, headers, node, vmid, is_lxc=False):
    ips = []
    if is_lxc:
        url = f"{base_url}/nodes/{node}/lxc/{vmid}/interfaces"
        ifaces = make_request("GET", url, headers, silent_agent=True) or []
        for iface in ifaces:
            if iface.get('name') == 'lo': continue
            if 'inet' in iface:
                ip = iface['inet'].split('/')[0]
                if ip != '127.0.0.1': ips.append(ip)
            if 'inet6' in iface:
                ip6 = iface['inet6'].split('/')[0]
                if ip6 != '::1' and not ip6.startswith('fe80'): ips.append(ip6)
    else:
        url = f"{base_url}/nodes/{node}/qemu/{vmid}/agent/network-get-interfaces"
        agent_data = make_request("GET", url, headers, silent_agent=True)
        if agent_data and 'result' in agent_data:
            for iface in agent_data['result']:
                if iface.get('name') == 'lo': continue
                for ip_info in iface.get('ip-addresses', []):
                    ip = ip_info.get('ip-address')
                    if ip and ip != '127.0.0.1' and ip != '::1' and not ip.startswith('fe80'):
                        ips.append(ip)
    
    return ", ".join(ips) if ips else "None"

def get_rrd_averages(base_url, headers, node, timeframe):
    rrd = make_request("GET", f"{base_url}/nodes/{node}/rrddata?timeframe={timeframe}", headers) or []
    cpu_sum, mem_sum, io_sum, swap_sum, count = 0, 0, 0, 0, 0
    
    for pt in rrd:
        if pt.get('cpu') is not None:  
            cpu_sum += pt['cpu'] * 100
            io_sum += pt.get('iowait', 0) * 100
            mem_tot = max(pt.get('memtotal', 1), 1) 
            mem_sum += (pt.get('memused', 0) / mem_tot) * 100
            swap_tot = max(pt.get('swaptotal', 1), 1) 
            swap_sum += (pt.get('swapused', 0) / swap_tot) * 100
            count += 1
            
    if count == 0: return 0.0, 0.0, 0.0, 0.0
    return round(cpu_sum/count, 2), round(mem_sum/count, 2), round(io_sum/count, 2), round(swap_sum/count, 2)

def write_bordered_row(ws, row, col, key, val, thin_border):
    c1 = ws.cell(row=row, column=col, value=key)
    c2 = ws.cell(row=row, column=col+1, value=val)
    c1.border = thin_border
    c2.border = thin_border

def main():
    servers = load_config("servers.json")
    zone_map = load_vlan_zones("vlan-zone.txt")
    
    all_instances_data = []
    all_nodes_data = []
    node_summaries = {}

    for server in servers:
        host = server['host']
        node = server['node']
        node_ip = host.split(':')[0]
        print(f"Fetching data from Node: {node} ({node_ip})...")
        node_summaries[node] = {"VM_Running": 0, "VM_Stopped": 0, "LXC_Running": 0, "LXC_Stopped": 0}
        
        base_url = f"https://{host}/api2/json"
        headers = {"Authorization": f"PVEAPIToken={server['token_id']}={server['token_secret']}"}
        
        # --- 1. FETCH NODE LEVEL INFORMATION ---
        node_status = make_request("GET", f"{base_url}/nodes/{node}/status", headers)
        if node_status:
            n_cpu_total = node_status.get('cpuinfo', {}).get('cpus', 0)
            n_ram_total = round(node_status.get('memory', {}).get('total', 0) / 1073741824, 2)
            n_model = node_status.get('cpuinfo', {}).get('model', 'Unknown CPU')
            n_sockets = node_status.get('cpuinfo', {}).get('sockets', 1)
            n_cpu_info = f"{n_cpu_total} x {n_model} ({n_sockets} Sockets)"
            n_kernel = node_status.get('kversion', 'Unknown')
            n_manager = node_status.get('pveversion', 'Unknown')
            n_boot = node_status.get('boot-info', {}).get('mode', 'Unknown')
            
            cur_cpu = round(node_status.get('cpu', 0) * 100, 2)
            cur_io = round(node_status.get('wait', 0) * 100, 2)
            
            n_mem_tot = max(node_status.get('memory', {}).get('total', 1), 1)
            cur_mem = round((node_status.get('memory', {}).get('used', 0) / n_mem_tot) * 100, 2)
            
            n_swap_tot = max(node_status.get('swap', {}).get('total', 1), 1)
            cur_swap = round((node_status.get('swap', {}).get('used', 0) / n_swap_tot) * 100, 2)
            
            hr_cpu, hr_mem, hr_io, hr_swap = get_rrd_averages(base_url, headers, node, 'hour')
            wk_cpu, wk_mem, wk_io, wk_swap = get_rrd_averages(base_url, headers, node, 'week')
            mo_cpu, mo_mem, mo_io, mo_swap = get_rrd_averages(base_url, headers, node, 'month')
            
            storages = make_request("GET", f"{base_url}/nodes/{node}/storage", headers) or []
            local_tot, local_used, shared_tot, shared_used = 0, 0, 0, 0
            for st in storages:
                if st.get('active') == 1 and st.get('total', 0) > 0:
                    if st.get('shared', 0) == 1:
                        shared_tot += st.get('total', 0)
                        shared_used += st.get('used', 0)
                    else:
                        local_tot += st.get('total', 0)
                        local_used += st.get('used', 0)
            
            local_total_gb = round(local_tot / 1073741824, 2)
            local_used_gb = round(local_used / 1073741824, 2)
            local_util_pct = round((local_used / local_tot * 100) if local_tot > 0 else 0, 2)
            shared_total_gb = round(shared_tot / 1073741824, 2)
            shared_used_gb = round(shared_used / 1073741824, 2)
            shared_util_pct = round((shared_used / shared_tot * 100) if shared_tot > 0 else 0, 2)

            all_nodes_data.append({
                "Node": node, "NodeIP": node_ip,
                "CPU_Total": n_cpu_total, "RAM_Total_GB": n_ram_total, "CPU_Info": n_cpu_info,
                "Kernel_Version": n_kernel, "Boot_Mode": n_boot, "Manager_Version": n_manager,
                "Local_Storage_Total_GB": local_total_gb, "Local_Storage_Used_GB": local_used_gb, "Local_Storage_Util_%": local_util_pct,
                "Shared_Storage_Total_GB": shared_total_gb, "Shared_Storage_Used_GB": shared_used_gb, "Shared_Storage_Util_%": shared_util_pct,
                "CPU_Cur_Min_%": cur_cpu, "Mem_Cur_Min_%": cur_mem, "IO_Cur_Min_%": cur_io, "Swap_Cur_Min_%": cur_swap,
                "CPU_Avg_Hr_%": hr_cpu, "Mem_Avg_Hr_%": hr_mem, "IO_Avg_Hr_%": hr_io, "Swap_Avg_Hr_%": hr_swap,
                "CPU_Avg_Wk_%": wk_cpu, "Mem_Avg_Wk_%": wk_mem, "IO_Avg_Wk_%": wk_io, "Swap_Avg_Wk_%": wk_swap,
                "CPU_Avg_Mo_%": mo_cpu, "Mem_Avg_Mo_%": mo_mem, "IO_Avg_Mo_%": mo_io, "Swap_Avg_Mo_%": mo_swap
            })
        
        # --- 2. FETCH VMs and LXCs ---
        vms = make_request("GET", f"{base_url}/nodes/{node}/qemu", headers) or []
        lxcs = make_request("GET", f"{base_url}/nodes/{node}/lxc", headers) or []
        
        if not vms and not lxcs:
            continue
            
        for vm in vms:
            vmid = int(vm.get('vmid'))
            name = vm.get('name', 'Unknown')
            status = vm.get('status', 'Unknown')
            uptime_sec = vm.get('uptime', 0)
            uptime_days = round(uptime_sec / 86400, 2) if uptime_sec else 0.0
            
            if status == "running": node_summaries[node]["VM_Running"] += 1
            else: node_summaries[node]["VM_Stopped"] += 1
            
            config = make_request("GET", f"{base_url}/nodes/{node}/qemu/{vmid}/config", headers) or {}
            vlan_tag = extract_vlan(config)
            zone_name = determine_zone(node, vlan_tag, zone_map)
            
            disks_str, hd_total_gb = extract_disks(config)
            total_cpu = int(config.get('sockets', 1)) * int(config.get('cores', 1))
            memory_gb = round(float(config.get('memory', vm.get('maxmem', 0) // 1048576)) / 1024, 2)
            notes = config.get('description', '')
            
            os_name = None
            ip_addresses = "None"
            if status == "running": 
                os_name = get_guest_os(base_url, headers, node, vmid)
                ip_addresses = get_guest_ips(base_url, headers, node, vmid, is_lxc=False)
            
            if not os_name: os_name = config.get('ostype', 'Unknown')
            
            all_instances_data.append({
                "Key": f"{node}_{vmid}", "Node": node, "NodeIP": node_ip, "VMID": vmid, 
                "Name": name, "State": status.upper(), "Uptime": uptime_days, "IP_Address": ip_addresses,
                "OS": os_name, "CPU": total_cpu, "Memory_GB": memory_gb, "VLAN": vlan_tag, "ZONE": zone_name,
                "Harddisks": disks_str, "HD_Total_GB": hd_total_gb, "Notes": notes
            })

        for lxc in lxcs:
            vmid = int(lxc.get('vmid'))
            name = lxc.get('name', 'Unknown')
            status = lxc.get('status', 'Unknown')
            uptime_sec = lxc.get('uptime', 0)
            uptime_days = round(uptime_sec / 86400, 2) if uptime_sec else 0.0
            
            if status == "running": node_summaries[node]["LXC_Running"] += 1
            else: node_summaries[node]["LXC_Stopped"] += 1
            
            config = make_request("GET", f"{base_url}/nodes/{node}/lxc/{vmid}/config", headers) or {}
            vlan_tag = extract_vlan(config)
            zone_name = determine_zone(node, vlan_tag, zone_map)
            
            disks_str, hd_total_gb = extract_disks(config)
            total_cpu = int(config.get('cores', lxc.get('cpus', 1)))
            memory_gb = round(float(config.get('memory', lxc.get('maxmem', 0) // 1048576)) / 1024, 2)
            notes = config.get('description', '')
            
            ip_addresses = "None"
            if status == "running":
                ip_addresses = get_guest_ips(base_url, headers, node, vmid, is_lxc=True)
            
            all_instances_data.append({
                "Key": f"{node}_{vmid}", "Node": node, "NodeIP": node_ip, "VMID": vmid, 
                "Name": name, "State": status.upper(), "Uptime": uptime_days, "IP_Address": ip_addresses,
                "OS": f"LXC Container ({config.get('ostype', 'Unknown')})", "CPU": total_cpu, 
                "Memory_GB": memory_gb, "VLAN": vlan_tag, "ZONE": zone_name, "Harddisks": disks_str, 
                "HD_Total_GB": hd_total_gb, "Notes": notes
            })

    # --- Prepare Export DataFrames ---
    df_vms = pd.DataFrame(all_instances_data)
    if not df_vms.empty:
        df_vms.insert(0, 'No', range(1, len(df_vms) + 1))
        # Insert ZONE column immediately after VLAN
        df_vms = df_vms[["No", "Key", "Node", "NodeIP", "VMID", "Name", "State", "Uptime", "IP_Address", "OS", "CPU", "Memory_GB", "VLAN", "ZONE", "Harddisks", "HD_Total_GB", "Notes"]]

    os.makedirs("result", exist_ok=True)
    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M")
    excel_filepath = os.path.join("result", f"vm_list_{timestamp}.xlsx")
    
    # --- Execute Excel Generation ---
    with pd.ExcelWriter(excel_filepath, engine='openpyxl') as writer:
        if not df_vms.empty:
            df_vms.to_excel(writer, index=False, sheet_name="VM Inventory")
            ws_vm = writer.sheets["VM Inventory"]
            
            grey_fill = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")
            state_col_idx = None
            
            for col_idx, cell in enumerate(ws_vm[1], start=1):
                if cell.value == "State":
                    state_col_idx = col_idx
                    break
            
            for row in ws_vm.iter_rows(min_row=2, max_row=ws_vm.max_row):
                if state_col_idx and row[state_col_idx - 1].value == "STOPPED":
                    for cell in row: cell.fill = grey_fill
            
            for col in ws_vm.columns:
                max_length = 0
                col_letter = col[0].column_letter
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length: max_length = len(str(cell.value))
                    except: pass
                ws_vm.column_dimensions[col_letter].width = min(max_length + 2, 60)

        if all_nodes_data:
            ws_node = writer.book.create_sheet("Node Information")
            thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
            center_align = Alignment(horizontal='center', vertical='center')
            
            ws_node.column_dimensions['A'].width = 15
            ws_node.column_dimensions['B'].width = 28
            ws_node.column_dimensions['C'].width = 65
            ws_node.column_dimensions['D'].width = 5
            ws_node.column_dimensions['E'].width = 28
            ws_node.column_dimensions['F'].width = 20
            
            r_idx = 1
            for i, n_data in enumerate(all_nodes_data, 1):
                ws_node.cell(row=r_idx, column=1, value=f"{i} {n_data['Node']}")
                r_idx += 1
                
                write_bordered_row(ws_node, r_idx,   2, 'CPU_Total', n_data['CPU_Total'], thin_border)
                write_bordered_row(ws_node, r_idx+1, 2, 'RAM_Total_GB', n_data['RAM_Total_GB'], thin_border)
                write_bordered_row(ws_node, r_idx+2, 2, 'CPU_Info', n_data['CPU_Info'], thin_border)
                write_bordered_row(ws_node, r_idx+3, 2, 'Kernel_Version', n_data['Kernel_Version'], thin_border)
                write_bordered_row(ws_node, r_idx+4, 2, 'Boot_Mode', n_data['Boot_Mode'], thin_border)
                write_bordered_row(ws_node, r_idx+5, 2, 'Manager_Version', n_data['Manager_Version'], thin_border)
                r_idx += 7
                
                write_bordered_row(ws_node, r_idx,   2, 'Local_Storage_Total_GB', n_data['Local_Storage_Total_GB'], thin_border)
                write_bordered_row(ws_node, r_idx+1, 2, 'Local_Storage_Used_GB', n_data['Local_Storage_Used_GB'], thin_border)
                write_bordered_row(ws_node, r_idx+2, 2, 'Local_Storage_Util_%', n_data['Local_Storage_Util_%'], thin_border)
                
                write_bordered_row(ws_node, r_idx,   5, 'Shared_Storage_Total_GB', n_data['Shared_Storage_Total_GB'], thin_border)
                write_bordered_row(ws_node, r_idx+1, 5, 'Shared_Storage_Used_GB', n_data['Shared_Storage_Used_GB'], thin_border)
                write_bordered_row(ws_node, r_idx+2, 5, 'Shared_Storage_Util_%', n_data['Shared_Storage_Util_%'], thin_border)
                r_idx += 4
                
                ws_node.cell(row=r_idx, column=2, value='Utilization')
                r_idx += 1
                
                timeframes = [
                    ('1 minute current', 'Cur_Min'), ('1 hour average', 'Avg_Hr'),
                    ('1 week average', 'Avg_Wk'), ('1 month average', 'Avg_Mo')
                ]
                
                for label, suffix in timeframes:
                    ws_node.cell(row=r_idx, column=2, value=label)
                    ws_node.merge_cells(start_row=r_idx, start_column=2, end_row=r_idx+3, end_column=2)
                    ws_node.cell(row=r_idx, column=2).alignment = center_align
                    
                    write_bordered_row(ws_node, r_idx,   3, f'CPU_{suffix}_%', n_data[f'CPU_{suffix}_%'], thin_border)
                    write_bordered_row(ws_node, r_idx+1, 3, f'Mem_{suffix}_%', n_data[f'Mem_{suffix}_%'], thin_border)
                    write_bordered_row(ws_node, r_idx+2, 3, f'IO_{suffix}_%', n_data[f'IO_{suffix}_%'], thin_border)
                    write_bordered_row(ws_node, r_idx+3, 3, f'Swap_{suffix}_%', n_data[f'Swap_{suffix}_%'], thin_border)
                    r_idx += 5
                    
                r_idx += 1

    print("\n" + "="*50)
    print("PROXMOX SUMMARY")
    print("="*50)
    for node_name, counts in node_summaries.items():
        print(f"Node: {node_name}")
        print(f"  - Total VM Running   : {counts['VM_Running']}")
        print(f"  - Total VM Stopped   : {counts['VM_Stopped']}")
        print(f"  - Total LXC Running  : {counts['LXC_Running']}")
        print(f"  - Total LXC Stopped  : {counts['LXC_Stopped']}")
        print("-" * 50)
            
    print(f"\n[Success] Formatted data exported to: {excel_filepath}")

if __name__ == "__main__":
    main()
